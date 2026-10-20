#!/usr/bin/env python3
"""Convert a Markdown file to a .docx with headings, tables, lists and inline formatting.

Usage: md2docx.py SOURCE.md OUTPUT.docx [brief]

The optional third argument selects the compact "brief" layout used for the
fixed-length JLC briefs (narrower margins, 9.5pt body, tight paragraph
spacing). The page geometry it applies is mirrored numerically in
pagecheck.py, which is what verifies that a one-page brief is one page.
Without the argument the default memo layout is used for everything else.
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

# Compact layout for the JLC briefs. Keep in sync with pagecheck.py BRIEF.
BRIEF_MARGIN_TB = 0.6      # inches, top and bottom
BRIEF_MARGIN_LR = 0.65     # inches, left and right
BRIEF_BODY_PT = 9.5        # body font size
BRIEF_LINE_PT = 11.5       # exact line height
BRIEF_SPACE_AFTER_PT = 3.0 # space after a body paragraph
BRIEF_TABLE_PT = 8.5       # table cell font size
# Heading level -> (size pt, space before pt, space after pt)
BRIEF_HEADINGS = {1: (15.0, 0.0, 4.0), 2: (11.5, 7.0, 3.0),
                  3: (10.0, 6.0, 2.0), 4: (9.5, 5.0, 2.0)}

INLINE = re.compile(
    r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|~~.+?~~|`[^`]+`|\*[^*]+?\*|\[[^\]]+\]\([^)]+\))",
    re.S,
)


def add_runs(par, text):
    """Add text to a paragraph, honouring bold/italic/strike/code/link markup."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        bold = italic = strike = code = False
        if tok.startswith("***") and tok.endswith("***") and len(tok) > 6:
            tok, bold, italic = tok[3:-3], True, True
        elif tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            tok, bold = tok[2:-2], True
        elif tok.startswith("~~") and tok.endswith("~~") and len(tok) > 4:
            tok, strike = tok[2:-2], True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            tok, code = tok[1:-1], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, italic = tok[1:-1], True
        else:
            m = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            if m:
                label, url = m.group(1), m.group(2)
                tok = f"{label} ({url})" if url.startswith("http") else label
        tok = re.sub(r"\\([\\`*_{}\[\]()#+\-.!=~|>])", r"\1", tok)
        run = par.add_run(tok)
        run.bold = bold
        run.italic = italic
        run.font.strike = strike
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x30, 0x30, 0x30)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_separator(line):
    return bool(re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", line)) and "-" in line


def add_table(doc, rows, brief=False):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(cols):
            text = row[j] if j < len(row) else ""
            par = cells[j].paragraphs[0]
            if brief:
                par.paragraph_format.space_after = Pt(1)
                par.paragraph_format.space_before = Pt(1)
            add_runs(par, text)
            for run in par.runs:
                if i == 0:
                    run.bold = True
                if brief:
                    run.font.size = Pt(BRIEF_TABLE_PT)
    if brief:
        # A sized spacer rather than a full empty body paragraph.
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.line_spacing = Pt(4)
        spacer.add_run("").font.size = Pt(4)
    else:
        doc.add_paragraph()


def add_rule(doc):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run("* * *")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x90, 0x90, 0x90)


def apply_brief_layout(doc):
    """Compact page geometry and typography for the fixed-length JLC briefs."""
    for section in doc.sections:
        section.top_margin = Inches(BRIEF_MARGIN_TB)
        section.bottom_margin = Inches(BRIEF_MARGIN_TB)
        section.left_margin = Inches(BRIEF_MARGIN_LR)
        section.right_margin = Inches(BRIEF_MARGIN_LR)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(BRIEF_BODY_PT)
    normal.paragraph_format.line_spacing = Pt(BRIEF_LINE_PT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(BRIEF_SPACE_AFTER_PT)

    for level, (size, before, after) in BRIEF_HEADINGS.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = Pt(size + 2)

    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = doc.styles[name]
        style.font.size = Pt(BRIEF_BODY_PT)
        style.paragraph_format.line_spacing = Pt(BRIEF_LINE_PT)
        style.paragraph_format.space_after = Pt(2)


def convert(src, dst, brief=False):
    lines = open(src, encoding="utf-8").read().replace("\r\n", "\n").split("\n")
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    if brief:
        apply_brief_layout(doc)

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if not line:
            i += 1
            continue

        # fenced code block
        if line.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            par = doc.add_paragraph()
            run = par.add_run("\n".join(buf))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue

        # table
        if line.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            rows = [split_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows, brief=brief)
            continue

        # horizontal rule
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", line):
            add_rule(doc)
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            head = doc.add_heading("", level=min(level, 4))
            add_runs(head, text)
            i += 1
            continue

        # blockquote
        if line.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            for qline in buf:
                if not qline.strip():
                    continue
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Pt(12 if brief else 24)
                par.paragraph_format.space_after = Pt(2 if brief else 4)
                add_runs(par, qline.strip())
                for run in par.runs:
                    run.font.size = Pt(BRIEF_BODY_PT if brief else 10)
            continue

        # bullet list
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", raw)
        if m:
            indent = len(m.group(1))
            par = doc.add_paragraph(style="List Bullet" if indent < 2 else "List Bullet 2")
            add_runs(par, m.group(2).strip())
            i += 1
            continue

        # numbered list
        m = re.match(r"^(\s*)\d+[.)]\s+(.*)$", raw)
        if m:
            indent = len(m.group(1))
            par = doc.add_paragraph(style="List Number" if indent < 2 else "List Number 2")
            add_runs(par, m.group(2).strip())
            i += 1
            continue

        # plain paragraph (merge wrapped lines)
        buf = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            s = nxt.strip()
            if (not s or s.startswith(("#", ">", "|", "```"))
                    or re.match(r"^\s*([-*+]\s+|\d+[.)]\s+)", nxt)
                    or re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", s)):
                break
            buf.append(s)
            i += 1
        par = doc.add_paragraph()
        add_runs(par, " ".join(buf))

    doc.save(dst)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2], brief=len(sys.argv) > 3 and sys.argv[3] == "brief")
    print("wrote", sys.argv[2])
